#!/usr/bin/env python3
"""Export analyzed PC-library reports into ONE styled wide-table .xlsx.

Standalone (CI/offline) and also invoked by the web panel's /api/export endpoint.
Reads each library's LATEST run report.json under <runs>/<lib>/<ts>/report.json and
flattens the business-facing fields into a single "分析汇总" sheet with a two-level
header (grouped 一级/二级 表头), 微软雅黑 font, merged group headers and light styling —
one row per library.

Usage:
  python3 scripts/export_xlsx.py --runs <runs_dir> --out <path.xlsx> [--names a,b,c]
"""
from __future__ import annotations

import argparse
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Single source of truth for dim-9 derivations (porting_class/adaptation_assessment/effort.level) and
# code_partition canonicalization — shared with assemble_report.py and web/server.js. Used to
# normalize a legacy (unstamped) report on load so the export matches the panel exactly.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import report_normalize  # noqa: E402

JOIN = "；"


# ── report loading (reused verbatim) ─────────────────────────────────────────

def _latest_report(lib_dir: str):
    """Return the parsed report.json of the newest run (ISO-timestamp dir) or None."""
    try:
        runs = sorted(
            d for d in os.listdir(lib_dir)
            if os.path.isfile(os.path.join(lib_dir, d, "report.json"))
        )
    except OSError:
        return None
    if not runs:
        return None
    path = os.path.join(lib_dir, runs[-1], "report.json")
    try:
        with open(path, encoding="utf-8") as fh:
            rep = json.load(fh)
    except (OSError, ValueError):
        return None
    # A report stamped at the CURRENT version is already the single-source-of-truth form; only
    # upgrade a legacy/unstamped/older-stamp one in memory so the export matches the panel exactly
    # (same version gate as web/server.js's /api/report serve path).
    stamp = rep.get("meta", {}).get("normalized_version") if isinstance(rep, dict) else None
    if isinstance(rep, dict) and (not isinstance(stamp, int) or stamp < report_normalize.NORMALIZED_VERSION):
        try:
            report_normalize.normalize_report(rep)
        except Exception:  # noqa: BLE001 — export must not fail on a normalize hiccup
            pass
    return rep


def _collect(runs_dir: str, names: set[str] | None):
    out = []
    try:
        libs = sorted(os.listdir(runs_dir))
    except OSError:
        return out
    for name in libs:
        lib_dir = os.path.join(runs_dir, name)
        if not os.path.isdir(lib_dir):
            continue
        if names and name not in names:
            continue
        rep = _latest_report(lib_dir)
        if rep is not None:
            out.append((name, rep))
    return out


# ── generic accessors (reused) ───────────────────────────────────────────────

def _g(obj, *path, default=None):
    """Safe nested get."""
    cur = obj
    for key in path:
        if isinstance(cur, dict):
            cur = cur.get(key)
        else:
            return default
        if cur is None:
            return default
    return cur


def _dep_container(r):
    """`dependencies` 正常是 schema dict {count, dependencies:[...]}，但部分模型运行
    直接产出裸数组。统一返回 dict 视图。"""
    d = r.get("dependencies")
    if isinstance(d, list):
        return {"dependencies": d}
    return d if isinstance(d, dict) else {}


def _join(items, key=None):
    if not items:
        return ""
    if key is None:
        return JOIN.join(str(x) for x in items if x not in (None, ""))
    vals = []
    for it in items:
        v = it.get(key) if isinstance(it, dict) else None
        if v not in (None, ""):
            vals.append(str(v))
    return JOIN.join(vals)


def _cell(v):
    """把模型形状漂移里混入的非标量值（dict/list）强制转成可写字符串，避免单个异常
    字段（如 build_env.build_system 偶发产出对象而非字符串）拖垮整包导出。"""
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (list, tuple)):
        return _join([_cell(x) for x in v])
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False, sort_keys=True)
    return str(v)


def _effort_days(ha: dict):
    e = (ha or {}).get("effort") or {}
    pd = e.get("person_days")
    if isinstance(pd, list) and len(pd) == 2:
        try:
            return [float(pd[0]), float(pd[1])]
        except (TypeError, ValueError):
            pass
    return None


def _cap_flag(r, key):
    for s in _g(r, "capability_profile", "scenarios", default=[]) or []:
        if s.get("key") == key and s.get("present"):
            return "是"
    return ""


# code_partition (dim 12) bucket LOC — field-name tolerance mirrors the panel's serve-time
# normalization (the agent occasionally emits total_loc/dir instead of loc/path).
def _bucket_loc(b):
    v = (b or {}).get("loc")
    if v is None:
        v = (b or {}).get("total_loc")
    if v is None:
        v = sum(int((m or {}).get("loc") or 0) for m in ((b or {}).get("modules") or []))
    return int(v or 0)


def _part_loc(r, cls):
    """Total LOC of one partition bucket class ('' when the block/bucket is absent)."""
    buckets = _g(r, "code_partition", "buckets", default=[]) or []
    hit = [_bucket_loc(b) for b in buckets if (b or {}).get("class") == cls]
    return sum(hit) if hit else ""


# ── label maps ───────────────────────────────────────────────────────────────

_LIC_CAT_LABELS = {"commercial": "商业协议", "strong_copyleft": "强传染协议",
                   "weak_copyleft": "弱传染协议", "permissive": "友好协议",
                   "undeclared": "未声明协议"}
_LEVEL_ZH = {"very_low": "极低", "low": "低", "medium": "中", "high": "高", "very_high": "极高"}
_PLAT_LABELS = {"windows": "Windows", "macos": "macOS", "linux": "Linux", "posix": "POSIX"}
_VENDOR_LABELS = {
    "google_firebase": "Firebase (Google)", "aws": "AWS", "gcp": "Google Cloud",
    "azure": "Azure", "alibaba_cloud": "阿里云", "tencent_cloud": "腾讯云",
    "huawei_cloud": "华为云", "supabase": "Supabase", "sentry": "Sentry",
    "cloudflare": "Cloudflare", "unknown": "未知厂商",
}
_CLOUD_CAT_LABELS = {
    "auth": "登录鉴权", "cloud_storage": "云存储", "database": "云数据库",
    "cloud_functions": "云函数", "push": "推送", "messaging": "消息",
    "analytics": "分析统计", "crash_reporting": "崩溃上报", "remote_config": "远程配置",
    "maps": "地图", "ml_ai": "AI 云推理", "ads": "广告", "hosting": "托管",
}
_KIND_LABELS = {"library": "库", "application": "应用", "framework": "框架",
                "tool": "工具", "cli": "命令行", "service": "服务", "plugin": "插件"}
_HSTATUS_LABELS = {"available": "可用", "partial": "部分", "unavailable": "不可用",
                   "unknown": "未知", "restricted": "受限"}
# 移植分级 5 档 effective_class（派生）——与 app.js TOPO_STATUS 一致
_PCLASS_LABELS = {"no_adaptation": "无需适配", "recompile_only": "仅交叉编译",
                  "needs_adaptation": "需适配·全部可适配",
                  "needs_adaptation_platform_partial": "需适配·平台差异有不可适配点",
                  "needs_adaptation_core_partial": "需适配·核心有不可适配点"}
# adaptation_assessment.overall — 是否可适配总判（派生，替代旧 feasibility）
_OVERALL_LABELS = {"adaptable": "可适配", "adaptable_with_tailoring": "可适配（部分平台特性需裁剪）",
                   "core_blocked": "核心功能不完全可适配"}
# functional_viability — 运行前提是否满足（目标侧派生轴，与 overall 正交）
_VIABILITY_LABELS = {"viable": "前提齐备", "viable_with_work": "有条件可用",
                     "blocked_external": "功能受阻·依赖外部前提", "unverified": "前提未核实"}

# ── 语义配色映射（列值 → _SEMANTIC_FILL 色键）─────────────────────────────────
# 与面板 app.js 的 TOPO_STATUS / *_CLS 语义一致；未知值 .get()→None → 不上色。
_PCLASS_FILL = {"no_adaptation": "blue", "recompile_only": "teal",
                "needs_adaptation": "amber",
                "needs_adaptation_platform_partial": "orange",
                "needs_adaptation_core_partial": "red"}
_OVERALL_FILL = {"adaptable": "green", "adaptable_with_tailoring": "amber",
                 "core_blocked": "red"}
_VIABILITY_FILL = {"viable": "green", "viable_with_work": "amber",
                   "blocked_external": "red", "unverified": "gray"}
_LEVEL_FILL = {"very_low": "green", "low": "green", "medium": "amber",
               "high": "orange", "very_high": "red"}
# target_assumptions.target_status —— 只对 required 项统计（元组顺序＝由严重到轻）
_TGT_ORDER = ("unavailable", "partial", "unknown")   # available 视为无阻碍
_TGT_FILL = {"unavailable": "red", "partial": "amber", "unknown": "gray"}
_TGT_MARK = {"unavailable": "不支持", "partial": "部分", "unknown": "未核实"}
# blockers.severity —— 先归一 high/medium/low → blocker/major/minor 再统计
_SEV_NORM = {"blocker": "blocker", "major": "major", "minor": "minor",
             "high": "blocker", "medium": "major", "low": "minor"}
_SEV_ORDER = ("blocker", "major", "minor")
_SEV_FILL = {"blocker": "red", "major": "orange", "minor": "gray"}
_SEV_MARK = {"blocker": "🔴阻塞", "major": "🟠主要", "minor": "⚪次要"}
# 富文本明细用：可适配性标签 + 目标状态 emoji/标签 + 圈号序号
_ADAPT_LABELS = {"adaptable": "可适配", "partial": "部分可适配", "unadaptable": "不可适配"}
_TGT_EMOJI = {"available": "🟢", "partial": "🟡", "unavailable": "🔴", "unknown": "⚪", "restricted": "🟠"}
_TGT_STATUS_LABEL = {"available": "已支持", "partial": "部分支持", "unavailable": "不支持",
                     "unknown": "未核实", "restricted": "受限"}
# target_status 严重度排序（越小越严重；per-item 排序用）
_TGT_SEV_RANK = {"unavailable": 0, "partial": 1, "restricted": 1, "unknown": 2, "available": 3}


def _circled(i):
    """①..⑳ 序号；超出范围回退 'N.'。"""
    return chr(0x245F + i) if 1 <= i <= 20 else f"{i}."


# ── per-column value functions ───────────────────────────────────────────────

def _v_name(name, r):
    return _g(r, "library", "name", default="") or name


def _v_source(name, r):
    url = _g(r, "library", "source_url", default="") or ""
    sub = _g(r, "library", "source_subpath", default="")
    return f"{url}（子目录: {sub}）" if sub else url


def _v_kind_form(name, r):
    """library.kind → 中文基名 + 综合推断的形态（GUI/命令行/服务），如「应用·GUI」。"""
    kind = _g(r, "library", "kind", default="library") or "library"
    base = _KIND_LABELS.get(kind, kind)
    types = set()
    for e in _g(r, "build_env", "entry_points", default=[]) or []:
        if isinstance(e, dict) and e.get("type"):
            types.add(e["type"])
        # entry_points 偶发混入裸字符串 —— 忽略即可
    gui = _cap_flag(r, "gui") == "是"
    forms = []
    if gui:
        forms.append("GUI")
    if "console_script" in types or kind == "cli":
        forms.append("命令行")
    if "service" in types or "framework_startup" in types or kind == "service":
        forms.append("服务")
    if not forms and ("launcher" in types or "main" in types):
        forms.append("GUI" if gui else "命令行")
    seen = []
    for f in forms:
        if f not in seen:
            seen.append(f)
    return base + ("·" + "/".join(seen) if seen else "")


def _v_desc(name, r):
    return _g(r, "function_summary", "summary", default="")


def _v_eco(name, r):
    eco = _g(r, "library", "ecosystem", default="") or ""
    bindings = _g(r, "library", "bindings", default=[]) or []
    extra = _join([b for b in bindings if b and b != eco])
    return f"{eco}（+{extra}）" if extra else eco


def _v_primary(name, r):
    return _g(r, "languages", "primary", default="")


def _v_license_name(name, r):
    return _g(r, "license", "name", default="") or _g(r, "license", "spdx", default="")


def _v_license_cat(name, r):
    c = _g(r, "license", "category", default="")
    return _LIC_CAT_LABELS.get(c, c or "")


def _v_dep_count(name, r):
    return len(_dep_container(r).get("dependencies", []) or [])


def _v_code(field):
    def fn(name, r):
        return _g(r, "code_metrics", field, "code")
    return fn


def _v_plat(key):
    def fn(name, r):
        v = _g(r, "code_metrics", "platform_adaptation", "by_platform", key, "code")
        return v if v not in (None, "") else 0
    return fn


def _v_part(cls):
    def fn(name, r):
        return _part_loc(r, cls)
    return fn


def _v_cap(key):
    def fn(name, r):
        for s in _g(r, "capability_profile", "scenarios", default=[]) or []:
            if s.get("key") == key and s.get("present"):
                st = _HSTATUS_LABELS.get(s.get("harmony_status"), s.get("harmony_status") or "")
                return f"是·{st}" if st else "是"
        return "—"
    return fn


def _v_cloud_vendors(name, r):
    seen = []
    for s in _g(r, "cloud_services", "services", default=[]) or []:
        v = s.get("vendor")
        if not v:
            continue
        label = _VENDOR_LABELS.get(v, v)
        if label not in seen:
            seen.append(label)
    return JOIN.join(seen) if seen else "—"


def _v_cloud_cats(name, r):
    seen = []
    for s in _g(r, "cloud_services", "services", default=[]) or []:
        for c in s.get("categories", []) or []:
            label = _CLOUD_CAT_LABELS.get(c, c)
            if label not in seen:
                seen.append(label)
    return "/".join(seen) if seen else "—"


def _v_pclass(name, r):
    # 移植分级 = 5 档派生 effective_class（回退到 3 值 porting_class）
    c = (_g(r, "harmony_adaptation", "adaptation_assessment", "effective_class", default="")
         or _g(r, "harmony_adaptation", "porting_class", default=""))
    return _Styled(_PCLASS_LABELS.get(c, c or ""), _PCLASS_FILL.get(c))


def _v_overall(name, r):
    o = _g(r, "harmony_adaptation", "adaptation_assessment", "overall", default="")
    return _Styled(_OVERALL_LABELS.get(o, o or ""), _OVERALL_FILL.get(o))


def _v_viability(name, r):
    v = _g(r, "harmony_adaptation", "functional_viability", default="")
    return _Styled(_VIABILITY_LABELS.get(v, v or ""), _VIABILITY_FILL.get(v))


def _v_level(name, r):
    lv = _g(r, "harmony_adaptation", "effort", "level", default="")
    return _Styled(_LEVEL_ZH.get(lv, lv or ""), _LEVEL_FILL.get(lv))


def _v_days(name, r):
    d = _effort_days(r.get("harmony_adaptation") or {})
    if not d:
        return ""
    lo, hi = d
    fmt = lambda x: (str(int(x)) if float(x).is_integer() else str(x))
    return f"{fmt(lo)}–{fmt(hi)}"


def _v_summary(name, r):
    return _g(r, "harmony_adaptation", "summary", default="")


def _v_target_assumptions(name, r):
    """目标平台能力假设 —— required 阻碍分布 headline + 逐条完整明细；色＝最严重 required。"""
    tas = [t for t in (_g(r, "harmony_adaptation", "target_assumptions", default=[]) or [])
           if isinstance(t, dict)]
    if not tas:
        return _Styled("—", None)
    req = [t for t in tas if t.get("required")]
    counts = {k: 0 for k in _TGT_ORDER}
    for t in req:
        st = t.get("target_status")
        if st in counts:
            counts[st] += 1
    head_parts = [f"{_TGT_EMOJI[k]}{_TGT_MARK[k]}{counts[k]}" for k in _TGT_ORDER if counts[k]]
    head = " ".join(head_parts) if head_parts else ("🟢全部已支持" if req else "")
    # required 优先、状态由重到轻（展示全部假设＝报告全集）
    ordered = sorted(tas, key=lambda t: (0 if t.get("required") else 1,
                                         _TGT_SEV_RANK.get(t.get("target_status"), 3)))
    items = []
    for i, t in enumerate(ordered, start=1):
        st = t.get("target_status") or "unknown"
        req_mark = "必需" if t.get("required") else "可选"
        lines = [f"{_circled(i)} {t.get('capability', '')} <{req_mark}> "
                 f"{_TGT_EMOJI.get(st, '⚪')}{_TGT_STATUS_LABEL.get(st, st)}"]
        if t.get("impact"):
            lines.append(f"   影响: {t['impact']}")
        if t.get("capability_key"):
            lines.append(f"   caps键: {t['capability_key']}")
        items.append("\n".join(lines))
    worst = next((k for k in _TGT_ORDER if counts[k]), None)
    fill = _TGT_FILL[worst] if worst else ("green" if req else None)
    body = (head + "\n──────\n" if head else "") + "\n".join(items)
    return _Styled(body, fill)


def _v_blockers(name, r):
    """移植阻碍点 —— 严重度分布 headline + 逐条完整明细（问题/改造/证据）；色＝最严重。"""
    bs = [b for b in (_g(r, "harmony_adaptation", "blockers", default=[]) or []) if isinstance(b, dict)]
    if not bs:
        return _Styled("无", "green")
    counts = {k: 0 for k in _SEV_ORDER}
    for b in bs:
        counts[_SEV_NORM.get(b.get("severity"), "major")] += 1
    head = " ".join(f"{_SEV_MARK[k]}{counts[k]}" for k in _SEV_ORDER if counts[k])
    bs.sort(key=lambda b: _SEV_ORDER.index(_SEV_NORM.get(b.get("severity"), "major")))
    items = []
    for i, b in enumerate(bs, start=1):
        sev = _SEV_NORM.get(b.get("severity"), "major")
        head1 = _SEV_MARK[sev]
        adapt = _ADAPT_LABELS.get(b.get("adaptability"))
        if adapt:
            head1 += "·" + adapt
        meta = [str(b[k]) for k in ("category", "source_dimension") if b.get(k)]
        if b.get("harmony_status"):
            meta.append("鸿蒙:" + str(b["harmony_status"]))
        lines = [f"{_circled(i)} {head1}" + (f"  [{'·'.join(meta)}]" if meta else "")]
        if b.get("issue"):
            lines.append(f"   问题: {b['issue']}")
        if b.get("remediation"):
            lines.append(f"   改造: {b['remediation']}")
        ev = [str(x) for x in (b.get("evidence") or []) if x]
        if ev:
            lines.append("   证据: " + "、".join(ev))
        items.append("\n".join(lines))
    worst = next(k for k in _SEV_ORDER if counts[k])
    return _Styled(head + "\n──────\n" + "\n".join(items), _SEV_FILL[worst])


def _v_critical_deps(name, r):
    """迁移关键路径依赖 —— 按 order 升序逐条（顺序/名称/人天/为何关键/关联）；中性无色。"""
    cds = [c for c in (_g(r, "harmony_adaptation", "critical_dependencies", default=[]) or [])
           if isinstance(c, dict) and c.get("name")]
    if not cds:
        return "—"

    def _key(ic):
        i, c = ic
        o = c.get("order")
        return (float(o) if isinstance(o, (int, float)) and o >= 1 else i + 1, i)

    ordered = [c for _, c in sorted(enumerate(cds), key=_key)]
    items = []
    for n, c in enumerate(ordered, start=1):
        o = c.get("order")
        num = int(o) if isinstance(o, (int, float)) and o >= 1 else n
        share = c.get("person_days_share")
        share_txt = ""
        if isinstance(share, list) and len(share) == 2:
            try:
                fmt = lambda x: (str(int(x)) if float(x).is_integer() else str(x))
                share_txt = f"  ({fmt(share[0])}–{fmt(share[1])} 人天)"
            except (TypeError, ValueError):
                share_txt = ""
        lines = [f"{num}. {c.get('name')}{share_txt}"]
        if c.get("why"):
            lines.append(f"   为何关键: {c['why']}")
        refs = [str(x) for x in (c.get("refs") or []) if x]
        if refs:
            lines.append(f"   关联: {' '.join(refs)}")
        items.append("\n".join(lines))
    return "\n".join(items)


# ── column spec (一级 → 二级) ────────────────────────────────────────────────
# Each entry: (top_title | None, [(sub_title, value_fn), ...]).
# top_title None → standalone column, header merged vertically across the two header rows.
GROUPS = [
    ("基本信息", [
        ("名称", _v_name), ("源码仓地址", _v_source), ("类型", _v_kind_form),
        ("描述", _v_desc), ("生态", _v_eco), ("主语言", _v_primary),
    ]),
    ("开源协议", [("协议类型", _v_license_name), ("协议友好类型", _v_license_cat)]),
    (None, [("依赖库数量", _v_dep_count)]),
    ("代码量(行)", [
        ("总体", _v_code("total")), ("生产", _v_code("production")),
        ("测试", _v_code("test")), ("样例", _v_code("example")),
    ]),
    ("平台适配代码量(行)", [
        ("Windows", _v_plat("windows")), ("macOS", _v_plat("macos")),
        ("Linux", _v_plat("linux")), ("POSIX", _v_plat("posix")),
    ]),
    ("代码分区(迁移复用性·工作量底座, LOC)", [
        ("直接复用", _v_part("reuse_direct")), ("重编译复用", _v_part("recompile_reuse")),
        ("需适配", _v_part("needs_adaptation")), ("无法适配", _v_part("unadaptable")),
    ]),
    ("能力画像", [
        ("GUI界面", _v_cap("gui")), ("3D渲染", _v_cap("rendering_3d")),
        ("媒体", _v_cap("media")), ("硬件/设备", _v_cap("hardware")),
    ]),
    ("云服务", [("厂商", _v_cloud_vendors), ("用途", _v_cloud_cats)]),
    ("鸿蒙适配评估", [
        ("移植分级", _v_pclass), ("代码适配", _v_overall), ("运行前提", _v_viability),
        ("难度", _v_level), ("工作量(人天)", _v_days), ("评估总结", _v_summary),
    ]),
    ("鸿蒙移植·阻碍/假设/关键路径", [
        ("目标能力假设(阻碍分布)", _v_target_assumptions),
        ("移植阻碍点(严重度分布)", _v_blockers),
        ("关键路径依赖(建议移植顺序)", _v_critical_deps),
    ]),
]

# per-column presentation overrides keyed by sub-title
WRAP_WIDTH = {"描述": 58, "评估总结": 58, "源码仓地址": 42, "用途": 22,
              "移植阻碍点(严重度分布)": 62, "目标能力假设(阻碍分布)": 50,
              "关键路径依赖(建议移植顺序)": 46}

# ── styling ──────────────────────────────────────────────────────────────────
FONT_NAME = "微软雅黑"
_TOP_FILL = PatternFill("solid", fgColor="B7C9E2")   # 一级表头（深）
_SUB_FILL = PatternFill("solid", fgColor="DDE6F0")    # 二级表头（浅）
_PART_TOP_FILL = PatternFill("solid", fgColor="E8D9B5")  # 代码分区组用暖色区分
_PART_SUB_FILL = PatternFill("solid", fgColor="F3ECD8")
_PORT_TOP_FILL = PatternFill("solid", fgColor="F0D9D2")  # 鸿蒙移植·阻碍组用淡红区分
_PORT_SUB_FILL = PatternFill("solid", fgColor="F7E9E4")
_ZEBRA_FILL = PatternFill("solid", fgColor="F7F9FC")
# 数据格语义软色（与面板 badge 底色一致）——由 _Styled.fill_key 选取
_SEMANTIC_FILL = {
    "green":  PatternFill("solid", fgColor="E3F6EA"),
    "teal":   PatternFill("solid", fgColor="DCEFEF"),
    "blue":   PatternFill("solid", fgColor="E4ECFB"),
    "amber":  PatternFill("solid", fgColor="FBF0D9"),
    "orange": PatternFill("solid", fgColor="FDEBD0"),
    "red":    PatternFill("solid", fgColor="FBE6E4"),
    "gray":   PatternFill("solid", fgColor="F0F3F7"),
}
_THIN = Side(style="thin", color="B0B8C4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_H_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CELL_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
_WRAP_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)


class _Styled(str):
    """带语义色键的单元格值；作为 str 子类，可原样流经 _cell() 与列宽测量。"""
    def __new__(cls, text, fill=None):
        s = super().__new__(cls, "" if text is None else str(text))
        s.fill_key = fill
        return s


def _flat_columns():
    """Flatten GROUPS into ordered [(sub_title, fn)] and remember group spans."""
    cols, spans, ci = [], [], 1
    for top, subs in GROUPS:
        start = ci
        for sub, fn in subs:
            cols.append((sub, fn))
            ci += 1
        spans.append((top, start, ci - 1))   # inclusive col range (1-based)
    return cols, spans


def build(entries, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "分析汇总"
    cols, spans = _flat_columns()
    ncols = len(cols)

    # header rows 1 (一级) + 2 (二级)
    for top, c1, c2 in spans:
        if top and top.startswith("代码分区"):
            top_fill, sub_fill = _PART_TOP_FILL, _PART_SUB_FILL
        elif top and top.startswith("鸿蒙移植"):
            top_fill, sub_fill = _PORT_TOP_FILL, _PORT_SUB_FILL
        else:
            top_fill, sub_fill = _TOP_FILL, _SUB_FILL
        if top is None:                       # standalone → vertical merge over both header rows
            ws.merge_cells(start_row=1, start_column=c1, end_row=2, end_column=c1)
            cell = ws.cell(row=1, column=c1, value=cols[c1 - 1][0])
            cell.fill = sub_fill
        else:
            if c2 > c1:
                ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
            tcell = ws.cell(row=1, column=c1, value=top)
            for c in range(c1, c2 + 1):
                ws.cell(row=1, column=c).fill = top_fill
            for c in range(c1, c2 + 1):
                scell = ws.cell(row=2, column=c, value=cols[c - 1][0])
                scell.fill = sub_fill
        # (data written below)

    # header styling
    for row in (1, 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(name=FONT_NAME, bold=True, size=11)
            cell.alignment = _H_ALIGN
            cell.border = _BORDER
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 22

    # data rows (start row 3)
    for ri, (name, rep) in enumerate(entries):
        r = 3 + ri
        for ci, (sub, fn) in enumerate(cols, start=1):
            try:
                val = _cell(fn(name, rep))
            except Exception as e:  # noqa: BLE001 — one bad field must not sink the row
                print(f"warn: {name} 列 {sub!r}: {e}", file=sys.stderr)
                val = ""
            cell = ws.cell(row=r, column=ci, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = _BORDER
            cell.alignment = _WRAP_ALIGN if sub in WRAP_WIDTH else _CELL_ALIGN
            fill_key = getattr(val, "fill_key", None)      # _Styled 携带的语义色键
            if fill_key in _SEMANTIC_FILL:
                cell.fill = _SEMANTIC_FILL[fill_key]
            elif ri % 2 == 1 and sub not in WRAP_WIDTH:
                cell.fill = _ZEBRA_FILL

    ws.freeze_panes = "C3"                     # freeze both header rows + 名称/源码仓 两列
    last_col = get_column_letter(ncols)
    ws.auto_filter.ref = f"A2:{last_col}{ws.max_row}"

    # column widths
    for ci, (sub, fn) in enumerate(cols, start=1):
        letter = get_column_letter(ci)
        if sub in WRAP_WIDTH:
            ws.column_dimensions[letter].width = WRAP_WIDTH[sub]
            continue
        longest = len(str(sub))
        # measure from the already-rendered cells (rows 3..)
        for r in range(3, ws.max_row + 1):
            v = ws.cell(row=r, column=ci).value
            if v is not None:
                longest = max(longest, min(len(str(v)), 24))
        ws.column_dimensions[letter].width = max(8, min(longest + 2, 26))

    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Export library reports to a styled wide-table .xlsx")
    ap.add_argument("--runs", required=True, help="runs/ directory")
    ap.add_argument("--out", required=True, help="output .xlsx path")
    ap.add_argument("--names", default="", help="comma-separated library names (default: all)")
    args = ap.parse_args(argv)

    names = {n.strip() for n in args.names.split(",") if n.strip()} or None
    entries = _collect(args.runs, names)
    if not entries:
        print("no analyzed libraries found", file=sys.stderr)
        return 2
    build(entries, args.out)
    print(f"exported {len(entries)} libraries -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
